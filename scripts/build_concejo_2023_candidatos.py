#!/usr/bin/env python3
"""
Genera src/data/electoral/concejo2023/<municipio>.json: resultados del CONCEJO 2023 desagregados por
partido y por candidato (voto preferente y voto solo por la lista) en los 125 municipios de Antioquia.

Fuente: libro "Concejo 2023 — Todos los resultados" (entregado por el equipo el 29-sep-2026), que reúne:
  (1) Escrutinio oficial, formularios E-24 CON / E-26 CON de la Registraduría Nacional del Estado Civil
      (escrutinios-2023.registraduria.gov.co), transcritos y verificados (lista + candidatos = total de
      cada partido): 13 municipios.
  (2) Preconteo oficial de la Registraduría (resultadosprec2023.registraduria.gov.co), solo donde el
      preconteo alcanzó al menos el 98 % de las mesas (el % exacto va en cada municipio): 86 municipios.
  26 municipios no tienen datos sólidos (el preconteo no llegó al 98 % de las mesas y no se transcribió el
  escrutinio): se generan con estado "sin-datos" y el % de mesas que alcanzó el preconteo. No se estima nada.

Crudo (no se redistribuye): _originales/registraduria/Concejo_Antioquia_2023_Todos_los_resultados.xlsx
Hojas: "Resultados" (una fila por candidato y por voto solo por la lista), "Totales por partido" y
"Resumen municipios" (habilitados, blanco, nulos, no marcados, estado del dato).

Controles (el script falla si alguno no se cumple):
  - los 125 municipios se cruzan con su código DANE (src/data/electoral/resultadosPuesto/indice.json);
  - por partido: voto solo lista + suma de candidatos = "Votos candidatos" + "Voto solo lista" de la hoja
    de totales;
  - por municipio: suma de los partidos = "Votos partidos" del resumen.

Uso: python scripts/build_concejo_2023_candidatos.py   (requiere openpyxl)
"""
from __future__ import annotations

import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / '_originales' / 'registraduria' / 'Concejo_Antioquia_2023_Todos_los_resultados.xlsx'
INDICE = ROOT / 'src' / 'data' / 'electoral' / 'resultadosPuesto' / 'indice.json'
OUT = ROOT / 'src' / 'data' / 'electoral' / 'concejo2023'

# Nombres del libro que no coinciden con el índice de Proteus
ALIAS = {
    'el carmen de viboral': 'carmen de viboral',
    'el penol': 'penol',
    'el retiro': 'retiro',
    'el santuario': 'santuario',
    'san pedro de los milagros': 'san pedro',
    'santa fe de antioquia': 'antioquia',
    'san vicente': 'san vicente ferrer',
}

# Municipios cuya cifra es oficial pero no representa la votación del municipio (se muestran con aviso)
INCOMPLETOS = {
    'Pueblorrico': 'La jornada del 29-oct-2023 no se completó: el E-26 CON del 06-nov-2023 solo recoge 103 votos '
                   'de 7.532 habilitados (el preconteo registró 0) y la Registraduría repitió la elección del '
                   'Concejo. Estas cifras no son la votación del municipio; el resultado de la repetición no está '
                   'en este libro.',
}

META = {
    'eleccion': 'Concejos municipales, elecciones territoriales del 29-oct-2023',
    'fuente': 'Registraduría Nacional del Estado Civil: escrutinio E-24 CON / E-26 CON (13 municipios) y '
              'preconteo oficial con al menos el 98 % de las mesas (86 municipios)',
    'nota': 'El preconteo al 100 % difiere del escrutinio en promedio 1 a 3 votos por candidato; el total '
            'municipal puede variar hasta ±5 % porque el escrutinio reclasifica votos. 26 municipios sin '
            'datos sólidos (preconteo < 98 % de mesas): no se muestran cifras.',
}


def norm(s: str) -> str:
    s = unicodedata.normalize('NFD', str(s).lower())
    return ' '.join(''.join(c for c in s if unicodedata.category(c) != 'Mn').replace('-', ' ').split())


def titulo(s: str) -> str:
    """Mayúscula inicial como en scripts/build_resultados_2023.mjs (la fuente viene en mayúsculas)."""
    t = re.sub(r'(^|[\s(\-"/])([a-záéíóúñü])', lambda m: m.group(1) + m.group(2).upper(), ' '.join(str(s).split()).lower())
    return re.sub(r'\b(De|Del|La|Las|Los|Y|E|En)\b', lambda m: m.group(1).lower(), t)


def pct_mesas(texto: str | None) -> float | None:
    m = re.search(r'([\d.,]+)\s*%', texto or '')
    return round(float(m.group(1).replace(',', '.')), 2) if m else None


def main() -> None:
    indice = json.loads(INDICE.read_text(encoding='utf-8'))
    por_nombre = {norm(v['nombre']): (k, v['2023']) for k, v in indice.items() if k.startswith('05')}

    wb = openpyxl.load_workbook(RAW, read_only=True, data_only=True)
    filas = list(wb['Resultados'].iter_rows(values_only=True))[1:]
    totales = list(wb['Totales por partido'].iter_rows(values_only=True))[1:]
    resumen = list(wb['Resumen municipios'].iter_rows(values_only=True))[1:]

    # municipio -> partido -> {soloLista, candidatos}
    datos: dict[str, dict[str, dict]] = defaultdict(dict)
    fuente_mpio: dict[str, str] = {}
    for mpio, partido, codigo, candidato, votos, fuente in filas:
        if codigo is None or votos is None:
            continue  # fila de aviso de los municipios sin datos sólidos
        p = datos[mpio].setdefault(partido, {'soloLista': 0, 'candidatos': []})
        if str(codigo) == '000':
            p['soloLista'] += int(votos)
        else:
            p['candidatos'].append({'codigo': str(codigo), 'nombre': titulo(candidato), 'votos': int(votos)})
        fuente_mpio.setdefault(mpio, fuente)

    # Control 1: totales por partido
    for mpio, partido, solo, cands, _total, _f in totales:
        if solo is None and cands is None:
            continue  # municipio sin datos sólidos
        p = datos[mpio][partido]
        if p['soloLista'] != int(solo) or sum(c['votos'] for c in p['candidatos']) != int(cands):
            raise SystemExit(f'Totales no cuadran: {mpio} / {partido}')

    OUT.mkdir(parents=True, exist_ok=True)
    vistos = set()
    resumen_estados = defaultdict(int)
    for mpio, estado, mesas, habilitados, votos_partidos, blanco, nulos, no_marcados, fuente in resumen:
        clave = ALIAS.get(norm(mpio), norm(mpio))
        if clave not in por_nombre:
            raise SystemExit(f'Municipio sin código DANE: {mpio}')
        dane, slug = por_nombre[clave]
        vistos.add(dane)
        solido = estado == 'Sólido'
        registro: dict = {
            'meta': META,
            'dane': dane,
            'municipio': mpio,
            'estado': ('incompleto' if mpio in INCOMPLETOS else 'solido') if solido else 'sin-datos',
            'pctMesas': 100.0 if 'escrutinio' in str(mesas) else pct_mesas(str(mesas)),
            'habilitados': int(habilitados) if habilitados is not None else None,
        }
        if solido:
            fuente = fuente or fuente_mpio.get(mpio, '')
            partidos = []
            for nombre, p in datos[mpio].items():
                cands = sorted(p['candidatos'], key=lambda c: (-c['votos'], c['codigo']))
                partidos.append({'nombre': titulo(nombre), 'total': p['soloLista'] + sum(c['votos'] for c in cands),
                                 'soloLista': p['soloLista'], 'candidatos': [[c['codigo'], c['nombre'], c['votos']] for c in cands]})
            partidos.sort(key=lambda p: -p['total'])
            suma = sum(p['total'] for p in partidos)
            if suma != int(votos_partidos):  # Control 2
                raise SystemExit(f'Suma de partidos {suma} != {votos_partidos} en {mpio}')
            registro.update({
                'tipo': 'escrutinio' if 'Escrutinio' in fuente else 'preconteo',
                'fuente': fuente,
                'votosPartidos': suma,
                'blanco': int(blanco or 0),
                'nulos': int(nulos) if nulos is not None else None,
                'noMarcados': int(no_marcados) if no_marcados is not None else None,
                'partidos': partidos,
            })
            if mpio in INCOMPLETOS:
                registro['nota'] = INCOMPLETOS[mpio]
            resumen_estados[registro['tipo'] if registro['estado'] == 'solido' else 'incompleto'] += 1
        else:
            if datos.get(mpio):
                raise SystemExit(f'{mpio} marcado sin datos pero trae filas')
            resumen_estados['sin-datos'] += 1
        (OUT / f'{slug}.json').write_text(json.dumps(registro, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')

    faltan = {dane for dane, _slug in por_nombre.values()} - vistos
    if faltan:
        raise SystemExit(f'Municipios de Antioquia sin fila en el libro: {sorted(faltan)}')
    print(f'{len(vistos)} municipios -> {OUT.relative_to(ROOT)}: {dict(resumen_estados)}')


if __name__ == '__main__':
    main()
