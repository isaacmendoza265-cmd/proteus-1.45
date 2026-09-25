"""Genera src/data/dane/antioquiaDane.json con datos oficiales del DANE para los 125 municipios.

Fuentes (en _originales/dane/, descargadas de dane.gov.co el 25-sep-2026):
- PPED-AreaMun-2018-2042_VP.xlsx: proyecciones de población municipal por área (actualizado 30-jul-2025).
- CNPV-2018-NBI.xlsx: Necesidades Básicas Insatisfechas, Censo Nacional de Población y Vivienda 2018.

Uso: python3 scripts/build_datos_dane.py  (requiere openpyxl)
"""
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / '_originales' / 'dane'
OUT = ROOT / 'src' / 'data' / 'dane' / 'antioquiaDane.json'
ANIO = 2026

pob = {}
wb = openpyxl.load_workbook(SRC / 'PPED-AreaMun-2018-2042_VP.xlsx', read_only=True, data_only=True)
for dp, _, mpio, nombre, anio, area, total in wb['PobMunicipalxÁrea'].iter_rows(min_row=10, max_col=7, values_only=True):
    if dp != '05' or anio != ANIO:
        continue
    r = pob.setdefault(mpio, {'nombre': nombre})
    key = {'Total': 'poblacion', 'Cabecera Municipal': 'poblacionCabecera', 'Centros Poblados y Rural Disperso': 'poblacionRural'}[area]
    r[key] = int(total)

wb = openpyxl.load_workbook(SRC / 'CNPV-2018-NBI.xlsx', read_only=True, data_only=True)
for row in wb['Municipios'].iter_rows(min_row=11, values_only=True):
    if row[0] != '05' or not row[2]:
        continue
    code = '05' + row[2]
    pob[code]['nbi2018'] = round(row[4], 2)
    pob[code]['miseria2018'] = round(row[5], 2)

assert len(pob) == 125 and all({'poblacion', 'nbi2018'} <= set(v) for v in pob.values()), 'faltan municipios'
out = {
    'meta': {
        'fuentePoblacion': f'DANE, proyecciones de población municipal por área 2018-2042 (actualización 30-jul-2025), año {ANIO}',
        'fuenteNbi': 'DANE, Necesidades Básicas Insatisfechas, CNPV 2018',
        'anioPoblacion': ANIO,
    },
    'municipios': dict(sorted(pob.items())),
}
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
print(f'OK {OUT.relative_to(ROOT)}: {len(pob)} municipios, población {ANIO} = {sum(v["poblacion"] for v in pob.values()):,}')
