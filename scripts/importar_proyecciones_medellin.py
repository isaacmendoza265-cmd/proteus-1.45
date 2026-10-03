#!/usr/bin/env python3
"""
Proyecciones de población de Medellín 2018-2030 del Distrito (DAP, Subdirección de Prospectiva, Información y
Evaluación Estratégica; actualización 2025, contrato interadministrativo con el DANE), por comuna y corregimiento y
por barrio y vereda, con sexo y grupos quinquenales de edad.

Fuente: https://www.medellin.gov.co/es/centro-documental/proyecciones-poblacion-viviendas-y-hogares/
Originales (descargados el 2-oct-2026 con permiso de Isaac) en _originales/medellin_proyecciones/:
  4-sexo-edad-quinquenal-comunas-2018-2030.xlsx   (una hoja por comuna/corregimiento y "Total Medellín")
  5-barrios-veredas-2018-2030.xlsx                (hoja "Barrios - Veredas - Sexo  Quin ")

Salidas:
  src/data/medellin/proyeccionesDistrito.json   estructura para Segmentos y la ficha (carga bajo demanda):
      {meta: {fuente, grupos: [17 quinquenales, último 80+], anios: [2018..2030]},
       unidades: {"medellin"|"comuna-N"|"med-correg-*"|"barrio-CCBB": {nombre, h: [[año x grupo]], m: [[...]]}}}
  src/data/motor/medellin-proyecciones-distrito.json   líneas para el motor de análisis (todas las herramientas de IA)

Uso: python3 scripts/importar_proyecciones_medellin.py
"""
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
ORIG = ROOT / '_originales/medellin_proyecciones'
ANIOS = list(range(2018, 2031))
GRUPOS = [f'{i}-{i + 4}' for i in range(0, 80, 5)] + ['80+']
FUENTE = ('Alcaldía de Medellín, Departamento Administrativo de Planeación: proyecciones de población 2018-2030 '
          '(actualización 2025, convenio con el DANE)')
CORREG = {'50': 'med-correg-palmitas', '60': 'med-correg-san-cristobal', '70': 'med-correg-altavista',
          '80': 'med-correg-san-antonio-de-prado', '90': 'med-correg-santa-elena'}
INDICE = json.load(open(ROOT / 'src/data/territorio/indiceTerritorios.json'))['medellin']


def norm(s):
    s = unicodedata.normalize('NFD', str(s)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', s)


def inicio_grupo(etq):
    """'0 a 4', '0-4', '80 a 84', '100 y más', '80+' -> edad inicial"""
    m = re.match(r'\s*(\d+)', str(etq))
    return int(m.group(1)) if m else None


def a_17(por_inicio):
    """{edad inicial: [13 años]} -> 17 grupos (80+ suma 80, 85, ..., 100)"""
    out = []
    for i in range(0, 80, 5):
        out.append(por_inicio.get(i, [0] * len(ANIOS)))
    resto = [sum(v[k] for e, v in por_inicio.items() if e >= 80) for k in range(len(ANIOS))]
    return out + [resto]


def transponer(grupos):  # [grupo][año] -> [año][grupo]
    return [[g[k] for g in grupos] for k in range(len(ANIOS))]


unidades = {}

# --- 1. Comunas, corregimientos y total (archivo 4) ---------------------------------------------------
wb = openpyxl.load_workbook(ORIG / '4-sexo-edad-quinquenal-comunas-2018-2030.xlsx', read_only=True, data_only=True)
por_nombre = {norm(v['nombre'].split(' - ')[-1].replace('Corregimiento ', '')): k for k, v in INDICE['divisiones'].items()}
for ws in wb.worksheets:
    if ws.title == 'Índice':
        continue
    datos = {'Hombres': {}, 'Mujeres': {}}
    for r in ws.iter_rows(values_only=True):
        v = [c for c in r if c is not None]
        if len(v) >= 15 and v[0] in datos and isinstance(v[2], (int, float)):
            e = inicio_grupo(v[1])
            if e is not None:
                datos[v[0]][e] = [int(x) for x in v[2:15]]
    if not datos['Hombres']:
        continue
    if ws.title.startswith('Total'):
        clave = 'medellin'
    else:
        clave = por_nombre.get(norm(ws.title))
        if not clave:  # nombres con variantes ("Laureles - Estadio", "Doce De Octubre")
            clave = next((k for n, k in por_nombre.items() if norm(ws.title) in n or n in norm(ws.title)), None)
    assert clave, f'hoja sin unidad: {ws.title}'
    unidades[clave] = {'nombre': ws.title, 'h': transponer(a_17(datos['Hombres'])), 'm': transponer(a_17(datos['Mujeres']))}

# --- 2. Barrios y veredas (archivo 5) ------------------------------------------------------------------
wb = openpyxl.load_workbook(ORIG / '5-barrios-veredas-2018-2030.xlsx', read_only=True, data_only=True)
ws = wb['Barrios - Veredas - Sexo  Quin ']
filas = {}
for r in ws.iter_rows(values_only=True):
    v = list(r)
    while v and v[0] is None:
        v.pop(0)
    if len(v) < 20 or not (isinstance(v[0], str) and v[0].isdigit()) or v[5] not in ('Hombres', 'Mujeres'):
        continue
    cod, nombre, area, sexo, edad = str(v[2]), v[3], v[4], v[5], v[6]
    e = inicio_grupo(edad)
    if e is None:
        continue
    filas.setdefault(cod, {'nombre': nombre, 'areas': {}})['areas'].setdefault(area, {'Hombres': {}, 'Mujeres': {}})[sexo][e] = [int(x or 0) for x in v[7:20]]
sin_indice = []
for cod, f in filas.items():
    area = f['areas'].get('Total') or next(iter(f['areas'].values()))
    clave = f'barrio-{cod}'
    if clave not in INDICE['subdivisiones']:
        sin_indice.append(f"{cod} {f['nombre']}")
    unidades[clave] = {'nombre': f['nombre'], 'h': transponer(a_17(area['Hombres'])), 'm': transponer(a_17(area['Mujeres']))}

meta = {'fuente': FUENTE, 'grupos': GRUPOS, 'anios': ANIOS,
        'nota': 'Actualización 2025 del Distrito con el DANE: el total de Medellín coincide con la proyección municipal del DANE (2.526.795 en 2026) y la desagrega por comuna, corregimiento, barrio y vereda. La serie vieja por comuna de observatorioComunas/populationData (2.650.662 en 2026) queda superada.'}
(ROOT / 'src/data/medellin').mkdir(exist_ok=True)
json.dump({'meta': meta, 'unidades': unidades}, open(ROOT / 'src/data/medellin/proyeccionesDistrito.json', 'w'), ensure_ascii=False, separators=(',', ':'))

# --- 3. Líneas para el motor de análisis -----------------------------------------------------------------
I26, I30, I18 = ANIOS.index(2026), ANIOS.index(2030), ANIOS.index(2018)
fmt = lambda n: f'{round(n):,}'.replace(',', '.')
pct = lambda a, b: f"{(100 * (b - a) / a):+.1f} %".replace('.', ',') if a else 'sin dato'


def tot(u, k):
    return sum(u['h'][k]) + sum(u['m'][k])


def rango(u, k, a, b):  # personas de a a b años (grupos quinquenales completos)
    return sum(u[s][k][i] for s in ('h', 'm') for i in range(len(GRUPOS)) if a <= i * 5 <= b)


def lineas(u):
    t18, t26, t30 = tot(u, I18), tot(u, I26), tot(u, I30)
    h26, m26 = sum(u['h'][I26]), sum(u['m'][I26])
    # 18 años o más: 2/5 del grupo 15-19 + 20 y más
    a = lambda k: 0.4 * (u['h'][k][3] + u['m'][k][3]) + sum(u[s][k][i] for s in ('h', 'm') for i in range(4, 17))
    j = lambda k: 0.4 * (u['h'][k][3] + u['m'][k][3]) + rango(u, k, 20, 25)
    may = lambda k: rango(u, k, 60, 80)
    return [
        f'Población proyectada 2026: {fmt(t26)} ({fmt(m26)} mujeres, {fmt(h26)} hombres). 2018: {fmt(t18)}. 2030: {fmt(t30)} ({pct(t26, t30)} frente a 2026).',
        f'Personas de 18 años o más: {fmt(a(I26))} en 2026 y {fmt(a(I30))} en 2030 ({pct(a(I26), a(I30))}).',
        f'Jóvenes de 18 a 29 años: {fmt(j(I26))} en 2026 y {fmt(j(I30))} en 2030 ({pct(j(I26), j(I30))}); 60 años o más: {fmt(may(I26))} en 2026 y {fmt(may(I30))} en 2030 ({pct(may(I26), may(I30))}).',
    ]


territorios = {k: {'lineas': lineas(u)} for k, u in unidades.items() if k != 'medellin'}
motor = {
    'meta': {'titulo': 'Proyección de población del Distrito de Medellín 2018-2030 (sexo y edad)', 'fuente': FUENTE,
             'nivel': 'oficial', 'categoria': 'población', 'corte': '2025',
             'nota': 'Desagrega por comuna, corregimiento, barrio y vereda la proyección del DANE (el total de Medellín coincide). Incluye la proyección a 2030'},
    'municipios': {'05001': {'lineas': ['Proyección del Distrito (total igual al del DANE). ' + x for x in lineas(unidades['medellin'])]}},
    'territorios': territorios,
}
json.dump(motor, open(ROOT / 'src/data/motor/medellin-proyecciones-distrito.json', 'w'), ensure_ascii=False, indent=0)
print(len(unidades), 'unidades;', 'barrios sin id en el índice:', len(sin_indice), sin_indice[:8])
print('Medellín 2026:', tot(unidades['medellin'], I26), ' 2030:', tot(unidades['medellin'], I30))
